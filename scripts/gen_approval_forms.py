# -*- coding: utf-8 -*-
"""
Step 1: 由「材料審批總表」生成 N 份分項材料報批表 xlsx（win32com 整張 copy 範本 + 只填值）。
關鍵：保留範本 logo / 中文字體 / 合併格，唔用 openpyxl clone。

用法：改下方 CONFIG，再跑：
  <WORKBUDDY_DIR>/binaries/python/envs/default/Scripts/python.exe gen_approval_forms.py
"""
import os
import openpyxl
import win32com.client as win32

# ===================== CONFIG（每個項目改呢度） =====================
BASE        = r"D:\Users\user\Desktop\2026-07-10-某學校-2025-2026學校年度某學校中學部暑期校舍工程"
SRC_SUMMARY = os.path.join(BASE, "材料審批", "材料審批總表-澳門某學校中學部暑期校舍工程_已填.xlsx")
SRC_TEMPLATE = os.path.join(BASE, "材料審批", "材料審批表-澳門某學校中學部暑期校舍工程.xlsx")
OUT         = os.path.join(BASE, "材料審批", "材料報批表-澳門某學校中學部暑期校舍工程_已填.xlsx")
# 總表欄位順序：1編號 2材料名 3BQ 4品牌 5數量 6單位 7貨期 8型號 9訂貨 10到貨 11備註 12備註
CODE_PREFIXES = ("AR-", "EL-", "AG-", "EG-", "AC-")
TICK = "■"
# 類別勾選格：建築/供排水/機電/電力
TRADE_CELLS = {"AR": "I7", "AG": "F8", "EG": "F8", "排風": "L7", "EL": "O7"}
# =====================================================================

def trade_cell(code, name):
    if code.startswith("AR"):
        return TRADE_CELLS["AR"]
    if code.startswith(("AG", "EG")):
        return TRADE_CELLS["AG"]
    if "排風" in str(name):
        return TRADE_CELLS["排風"]
    return TRADE_CELLS["EL"]

# 讀數據（openpyxl 只讀）
wb = openpyxl.load_workbook(SRC_SUMMARY, data_only=True)
ws = wb.active
items = []
for r in range(1, ws.max_row + 1):
    a = ws.cell(row=r, column=1).value
    if not a:
        continue
    a = str(a).strip()
    if a.startswith(CODE_PREFIXES):
        items.append({
            "code":  a,
            "name":  ws.cell(row=r, column=2).value or "",
            "bq":    ws.cell(row=r, column=3).value or "",
            "brand": ws.cell(row=r, column=4).value or "",
            "qty":   ws.cell(row=r, column=5).value,
            "unit":  ws.cell(row=r, column=6).value or "",
            "lead":  ws.cell(row=r, column=7).value or "",
            "model": ws.cell(row=r, column=8).value or "",
            "note":  ws.cell(row=r, column=12).value or "",
        })
wb.close()
print("items:", len(items))

excel = win32.Dispatch("Excel.Application")
excel.Visible = False
excel.DisplayAlerts = False
try:
    wb_t = excel.Workbooks.Open(SRC_TEMPLATE, ReadOnly=True)
    ws_t = wb_t.Worksheets(1)
    wb_out = excel.Workbooks.Add()
    for it in items:
        n_before = wb_out.Worksheets.Count
        ws_t.Copy(After=wb_out.Worksheets(n_before))   # 整張 copy：圖片/字體/格式 1:1
        sh = wb_out.Worksheets(wb_out.Worksheets.Count)
        sh.Range("D5").Value = it["code"]
        try:
            sh.Range("H5:R5").Merge()
        except Exception:
            pass
        sh.Range("H5").Value = it["bq"]
        sh.Range("D7").Value = it["name"]
        sh.Range(trade_cell(it["code"], it["name"])).Value = TICK
        sh.Range("D9").Value = it["brand"]
        sh.Range("F10").Value = TICK
        qty_txt = f"{it['qty']} {it['unit']}" if it["unit"] else str(it["qty"])
        sh.Range("C17").Value = qty_txt
        if it["lead"]:
            sh.Range("J17").Value = str(it["lead"])
        note = it["note"] or ""
        if not note and it["model"]:
            note = f"型號：{it['model']}"
        if note:
            sh.Range("C19").Value = note
        try:
            sh.Name = it["code"]
        except Exception:
            pass
    for s in list(wb_out.Worksheets):
        if s.Name == "Sheet1":
            s.Delete()
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb_out.SaveAs(OUT, FileFormat=51)   # xlOpenXMLWorkbook
    wb_out.Close(SaveChanges=False)
    wb_t.Close(SaveChanges=False)
    print("SAVED:", OUT)
finally:
    excel.Quit()
print("DONE")
